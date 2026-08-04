#!/usr/bin/env python3
"""Sinh file Word (.docx) báo cáo dự án STS từ nội dung có cấu trúc.

Chạy: python3 docs/build_docx.py
Kết quả: docs/bao-cao-du-an-sts.docx
"""

import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
LIGHT = RGBColor(0xD9, 0xE2, 0xF3)


def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)
        set_cell_background(hdr[i], "1F4E79")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(10)
    if widths:
        for i, w in enumerate(widths):
            for r in table.rows:
                r.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = ACCENT
    return h


def add_bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def add_numbered(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def add_code_block(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fldChar = OxmlElement("w:fldChar")
    fldChar.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "Nhấn Ctrl+Click / cập nhật trường (F9) để hiển thị mục lục."
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(t)
    run._r.append(fldChar3)


def add_page_number_footer(doc):
    section = doc.sections[-1]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Trang ")
    run.font.size = Pt(9)
    fld1 = OxmlElement("w:fldSimple")
    fld1.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld1)


def build():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ---------------- Trang bìa ----------------
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO DỰ ÁN")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = ACCENT

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HỆ THỐNG QUẢN LÝ ĐIỂM RÈN LUYỆN SINH VIÊN")
    r.bold = True
    r.font.size = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Student Training Score (STS)")
    r.italic = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    for _ in range(6):
        doc.add_paragraph()

    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Tên dự án", "Hệ thống Quản lý Điểm Rèn luyện Sinh viên (STS)"),
        ("Phiên bản tài liệu", "1.0"),
        ("Ngày lập", "04/08/2026"),
        ("Trạng thái", "Bản nháp"),
    ]
    for i, (k, v) in enumerate(info_data):
        c0 = info.rows[i].cells[0]
        c0.text = ""
        rr = c0.paragraphs[0].add_run(k)
        rr.bold = True
        c0.width = Inches(2.2)
        set_cell_background(c0, "D9E2F3")
        c1 = info.rows[i].cells[1]
        c1.text = v
        c1.width = Inches(4.0)

    doc.add_page_break()

    # ---------------- Mục lục ----------------
    add_heading(doc, "Mục lục", 1)
    add_toc(doc)
    doc.add_page_break()

    # ---------------- 1. Tóm tắt ----------------
    add_heading(doc, "1. Tóm tắt", 1)
    doc.add_paragraph(
        "Dự án STS (Student Training Score) xây dựng hệ thống quản lý điểm rèn luyện "
        "sinh viên, hỗ trợ nhà trường theo dõi, đánh giá và tổng hợp kết quả rèn luyện "
        "theo học kỳ/năm học một cách minh bạch, thống nhất và dễ kiểm soát."
    )
    doc.add_paragraph("Hệ thống hướng tới các mục tiêu:")
    add_bullets(doc, [
        "Số hóa quy trình đánh giá điểm rèn luyện.",
        "Giảm thao tác thủ công, sai sót và trùng lặp dữ liệu.",
        "Tăng tính minh bạch giữa sinh viên, cố vấn học tập và phòng công tác sinh viên.",
        "Cung cấp báo cáo, thống kê phục vụ quản lý đào tạo.",
    ])

    # ---------------- 2. Giới thiệu ----------------
    add_heading(doc, "2. Giới thiệu", 1)
    add_heading(doc, "2.1. Bối cảnh", 2)
    doc.add_paragraph(
        "Điểm rèn luyện là chỉ số quan trọng phản ánh ý thức, thái độ và kết quả tham gia "
        "hoạt động của sinh viên. Quy trình đánh giá thường gồm nhiều bước: tự đánh giá của "
        "sinh viên, xác nhận của lớp/cố vấn, và phê duyệt của khoa/phòng CTSV."
    )
    doc.add_paragraph("Việc thực hiện thủ công bằng Excel/giấy hiện gây ra:")
    add_bullets(doc, [
        "Khó đối chiếu và truy vết lịch sử chỉnh sửa.",
        "Tốn thời gian tổng hợp theo lớp, khoa, toàn trường.",
        "Thiếu kênh phản hồi rõ ràng khi sinh viên khiếu nại điểm.",
    ])

    add_heading(doc, "2.2. Mục tiêu dự án", 2)
    add_table(doc,
              ["STT", "Mục tiêu", "Mô tả"],
              [
                  ["1", "Quản lý tập trung", "Lưu trữ điểm rèn luyện theo học kỳ, năm học, lớp, khoa"],
                  ["2", "Phân quyền rõ ràng", "Sinh viên, cố vấn, cán bộ khoa, admin có vai trò riêng"],
                  ["3", "Quy trình chuẩn hóa", "Tự đánh giá → xác nhận → phê duyệt"],
                  ["4", "Báo cáo & thống kê", "Xuất báo cáo theo lớp/khoa/toàn trường"],
                  ["5", "Minh bạch", "Sinh viên xem được điểm và lý do điều chỉnh"],
              ],
              widths=[0.6, 2.0, 4.0])

    add_heading(doc, "2.3. Phạm vi", 2)
    doc.add_paragraph("Trong phạm vi (MVP):", style="Intense Quote")
    add_bullets(doc, [
        "Đăng nhập / phân quyền người dùng.",
        "Quản lý danh mục: khoa, lớp, năm học, học kỳ, tiêu chí đánh giá.",
        "Sinh viên tự chấm / gửi phiếu đánh giá.",
        "Cố vấn xác nhận / điều chỉnh điểm.",
        "Khoa / CTSV phê duyệt.",
        "Xem lịch sử và xuất báo cáo cơ bản (CSV/Excel/PDF).",
    ])
    doc.add_paragraph("Ngoài phạm vi (giai đoạn sau):", style="Intense Quote")
    add_bullets(doc, [
        "Tích hợp SSO trường / LMS.",
        "App mobile chuyên biệt.",
        "AI gợi ý điểm / phát hiện bất thường nâng cao.",
    ])

    # ---------------- 3. Phân tích yêu cầu ----------------
    add_heading(doc, "3. Phân tích yêu cầu", 1)
    add_heading(doc, "3.1. Đối tượng sử dụng (Actors)", 2)
    add_table(doc,
              ["Vai trò", "Mô tả", "Nhu cầu chính"],
              [
                  ["Sinh viên", "Người được đánh giá", "Tự đánh giá, xem điểm, khiếu nại"],
                  ["Cố vấn học tập / GVCN", "Quản lý lớp", "Xác nhận, điều chỉnh, nhận xét"],
                  ["Cán bộ khoa", "Quản lý cấp khoa", "Phê duyệt, thống kê theo khoa"],
                  ["Phòng CTSV / Admin", "Quản trị hệ thống", "Cấu hình tiêu chí, phê duyệt cuối, báo cáo"],
                  ["Lãnh đạo", "Người ra quyết định", "Xem dashboard, xuất báo cáo tổng hợp"],
              ],
              widths=[1.8, 1.8, 2.8])

    add_heading(doc, "3.2. Yêu cầu chức năng", 2)
    add_numbered(doc, [
        "Xác thực & phân quyền: đăng nhập, phân role, giới hạn thao tác theo quyền.",
        "Quản lý danh mục: khoa, ngành, lớp, sinh viên, năm học, học kỳ.",
        "Cấu hình khung điểm rèn luyện: nhóm tiêu chí, điểm tối đa, quy tắc cộng/trừ.",
        "Phiếu tự đánh giá: sinh viên nhập điểm theo tiêu chí, đính kèm minh chứng.",
        "Quy trình phê duyệt nhiều cấp: cố vấn → khoa → CTSV.",
        "Điều chỉnh điểm có lý do: ghi nhận người sửa, thời gian, nội dung thay đổi.",
        "Thống kê & báo cáo: theo lớp, khoa, xếp loại.",
        "Thông báo: nhắc hạn nộp, kết quả phê duyệt, yêu cầu bổ sung minh chứng.",
    ])

    add_heading(doc, "3.3. Yêu cầu phi chức năng", 2)
    add_table(doc,
              ["Nhóm", "Yêu cầu"],
              [
                  ["Hiệu năng", "Truy vấn danh sách lớp ~200 SV dưới 2 giây"],
                  ["Bảo mật", "Mã hóa mật khẩu, phân quyền chặt, ghi audit log"],
                  ["Khả dụng", "Giao diện tiếng Việt, responsive desktop/mobile"],
                  ["Độ tin cậy", "Sao lưu định kỳ, khôi phục dữ liệu khi sự cố"],
                  ["Khả mở rộng", "Dễ thêm tiêu chí / cấp phê duyệt mới"],
              ],
              widths=[1.6, 4.8])

    # ---------------- 4. Thiết kế hệ thống ----------------
    add_heading(doc, "4. Thiết kế hệ thống", 1)
    add_heading(doc, "4.1. Kiến trúc tổng quan", 2)
    add_code_block(doc,
                   "Web Client (Browser)  <-->  API Backend (REST / Auth)  <-->  Database (PostgreSQL)\n"
                   "                                     |\n"
                   "                                     v\n"
                   "                          File Storage (minh chứng)")

    add_heading(doc, "4.2. Luồng nghiệp vụ chính", 2)
    add_code_block(doc,
                   "Sinh viên tự đánh giá\n"
                   "        -> Cố vấn xác nhận / điều chỉnh\n"
                   "        -> Khoa phê duyệt\n"
                   "        -> CTSV / Admin chốt điểm\n"
                   "        -> Công bố kết quả + xuất báo cáo")

    add_heading(doc, "4.3. Phân quyền thao tác (tóm tắt)", 2)
    add_table(doc,
              ["Chức năng", "SV", "Cố vấn", "Khoa", "Admin"],
              [
                  ["Tự đánh giá", "x", "", "", ""],
                  ["Xác nhận lớp", "", "x", "", ""],
                  ["Phê duyệt khoa", "", "", "x", ""],
                  ["Cấu hình tiêu chí", "", "", "", "x"],
                  ["Xuất báo cáo toàn trường", "", "", "x", "x"],
              ],
              widths=[2.6, 0.9, 0.9, 0.9, 0.9])

    # ---------------- 5. Mô hình dữ liệu ----------------
    add_heading(doc, "5. Mô hình dữ liệu", 1)
    add_heading(doc, "5.1. Các thực thể chính", 2)
    add_table(doc,
              ["Thực thể", "Mô tả"],
              [
                  ["User", "Tài khoản người dùng (role, trạng thái)"],
                  ["Student", "Hồ sơ sinh viên (MSSV, lớp, khoa)"],
                  ["Faculty / Class", "Khoa, lớp học"],
                  ["AcademicTerm", "Năm học / học kỳ"],
                  ["CriteriaGroup / Criterion", "Nhóm tiêu chí và tiêu chí chấm điểm"],
                  ["EvaluationSheet", "Phiếu đánh giá của 1 SV trong 1 học kỳ"],
                  ["EvaluationItem", "Điểm từng tiêu chí trong phiếu"],
                  ["ApprovalStep", "Các bước phê duyệt / trạng thái"],
                  ["Evidence", "File minh chứng đính kèm"],
                  ["AuditLog", "Lịch sử thay đổi"],
              ],
              widths=[2.4, 4.0])

    add_heading(doc, "5.2. Trạng thái phiếu đánh giá", 2)
    add_table(doc,
              ["Trạng thái", "Ý nghĩa"],
              [
                  ["DRAFT", "Sinh viên đang soạn"],
                  ["SUBMITTED", "Đã gửi, chờ cố vấn"],
                  ["ADVISOR_REVIEWED", "Cố vấn đã xử lý"],
                  ["FACULTY_APPROVED", "Khoa đã duyệt"],
                  ["FINALIZED", "CTSV chốt điểm"],
                  ["REJECTED", "Trả lại để chỉnh sửa"],
              ],
              widths=[2.4, 4.0])

    add_heading(doc, "5.3. Xếp loại điểm rèn luyện (tham chiếu)", 2)
    add_table(doc,
              ["Khoảng điểm", "Xếp loại"],
              [
                  ["90 – 100", "Xuất sắc"],
                  ["80 – <90", "Tốt"],
                  ["65 – <80", "Khá"],
                  ["50 – <65", "Trung bình"],
                  ["<50", "Yếu"],
              ],
              widths=[2.4, 4.0])
    note = doc.add_paragraph()
    nr = note.add_run("Lưu ý: khung điểm/xếp loại cần cấu hình theo quy chế thực tế của từng trường.")
    nr.italic = True

    # ---------------- 6. Chức năng chính ----------------
    add_heading(doc, "6. Chức năng chính", 1)
    add_heading(doc, "6.1. Module quản trị", 2)
    add_bullets(doc, [
        "Quản lý người dùng, gán vai trò.",
        "Import danh sách sinh viên (CSV/Excel).",
        "Cấu hình năm học, học kỳ, khung tiêu chí.",
        "Mở / khóa đợt đánh giá.",
    ])
    add_heading(doc, "6.2. Module sinh viên", 2)
    add_bullets(doc, [
        "Xem khung tiêu chí và hướng dẫn chấm.",
        "Điền phiếu tự đánh giá, upload minh chứng.",
        "Theo dõi trạng thái phê duyệt.",
        "Xem điểm cuối và lịch sử điều chỉnh.",
    ])
    add_heading(doc, "6.3. Module cố vấn / khoa / CTSV", 2)
    add_bullets(doc, [
        "Duyệt danh sách phiếu theo lớp/khoa.",
        "Điều chỉnh điểm kèm lý do.",
        "Phê duyệt hàng loạt.",
        "Thống kê tỷ lệ xếp loại, sinh viên chưa nộp.",
    ])
    add_heading(doc, "6.4. Module báo cáo", 2)
    add_bullets(doc, [
        "Báo cáo theo lớp / khoa / toàn trường.",
        "Danh sách xếp loại Xuất sắc → Yếu.",
        "Xuất file Excel / PDF phục vụ lưu trữ và họp xét.",
    ])

    # ---------------- 7. Công nghệ ----------------
    add_heading(doc, "7. Công nghệ đề xuất", 1)
    add_table(doc,
              ["Lớp", "Đề xuất", "Lý do"],
              [
                  ["Frontend", "React / Next.js", "UI hiện đại, dễ bảo trì"],
                  ["Backend", "NestJS hoặc Spring Boot", "API rõ ràng, phân quyền tốt"],
                  ["Database", "PostgreSQL", "Đáng tin cậy, hỗ trợ quan hệ phức tạp"],
                  ["Auth", "JWT + refresh token (hoặc SSO)", "Bảo mật, mở rộng được"],
                  ["File", "Object storage (S3-compatible)", "Lưu minh chứng"],
                  ["Deploy", "Docker + CI/CD", "Triển khai ổn định"],
              ],
              widths=[1.4, 2.6, 2.4])

    # ---------------- 8. Kế hoạch ----------------
    add_heading(doc, "8. Kế hoạch triển khai", 1)
    add_table(doc,
              ["Giai đoạn", "Nội dung chính", "Kết quả bàn giao"],
              [
                  ["Giai đoạn 1", "Khảo sát quy chế, chốt yêu cầu MVP", "Tài liệu SRS"],
                  ["Giai đoạn 2", "Thiết kế DB, API, UI wireframe", "Bản thiết kế kỹ thuật"],
                  ["Giai đoạn 3", "Phát triển core: auth, danh mục, phiếu", "Bản demo nội bộ"],
                  ["Giai đoạn 4", "Phê duyệt nhiều cấp, báo cáo, audit log", "Bản UAT"],
                  ["Giai đoạn 5", "Kiểm thử, sửa lỗi, triển khai production", "Hệ thống chính thức"],
                  ["Giai đoạn 6", "Vận hành, thu thập phản hồi, tối ưu", "Báo cáo nghiệm thu"],
              ],
              widths=[1.4, 3.0, 2.0])

    # ---------------- 9. Rủi ro ----------------
    add_heading(doc, "9. Rủi ro và biện pháp", 1)
    add_table(doc,
              ["Rủi ro", "Mức độ", "Biện pháp"],
              [
                  ["Quy chế điểm thay đổi theo năm", "Cao", "Thiết kế tiêu chí cấu hình được, không hard-code"],
                  ["Dữ liệu import sai từ Excel", "Trung bình", "Validate chặt, cho phép rollback theo đợt"],
                  ["Quá tải cuối đợt đánh giá", "Trung bình", "Cache danh mục, tối ưu truy vấn, giới hạn upload"],
                  ["Tranh chấp điểm / khiếu nại", "Cao", "Lưu audit log đầy đủ, bắt buộc nhập lý do sửa"],
                  ["Người dùng chưa quen quy trình số", "Trung bình", "Hướng dẫn trong app, video ngắn, hỗ trợ lớp"],
              ],
              widths=[2.4, 1.2, 2.8])

    # ---------------- 10. Kết luận ----------------
    add_heading(doc, "10. Kết luận và hướng phát triển", 1)
    doc.add_paragraph(
        "Dự án STS hướng tới số hóa toàn bộ vòng đời đánh giá điểm rèn luyện sinh viên: từ "
        "tự đánh giá, xác nhận nhiều cấp, đến báo cáo và lưu trữ lịch sử. Khi hoàn thành MVP, "
        "hệ thống sẽ giúp nhà trường giảm công sức vận hành, tăng minh bạch và tạo nền tảng "
        "dữ liệu cho công tác sinh viên."
    )
    doc.add_paragraph("Hướng phát triển tiếp theo:")
    add_numbered(doc, [
        "Tích hợp hệ thống thông tin sinh viên / SSO của trường.",
        "Thông báo real-time (email / Zalo OA / app).",
        "Dashboard phân tích xu hướng rèn luyện theo khóa, ngành.",
        "Module khiếu nại – phúc khảo online.",
        "Ứng dụng di động cho sinh viên và cố vấn.",
    ])

    # ---------------- 11. Phụ lục ----------------
    add_heading(doc, "11. Phụ lục", 1)
    add_heading(doc, "A. Thuật ngữ", 2)
    add_table(doc,
              ["Thuật ngữ", "Giải thích"],
              [
                  ["Điểm rèn luyện", "Điểm đánh giá ý thức, thái độ, hoạt động của sinh viên"],
                  ["Phiếu đánh giá", "Bản ghi điểm theo tiêu chí của 1 SV trong 1 học kỳ"],
                  ["Cố vấn học tập", "Giảng viên phụ trách lớp, xác nhận điểm cấp lớp"],
                  ["CTSV", "Phòng Công tác Sinh viên"],
                  ["MVP", "Minimum Viable Product – phiên bản tối thiểu có thể dùng"],
              ],
              widths=[2.0, 4.4])

    add_heading(doc, "B. Lịch sử tài liệu", 2)
    add_table(doc,
              ["Phiên bản", "Ngày", "Người lập", "Nội dung"],
              [
                  ["1.0", "04/08/2026", "—", "Khởi tạo báo cáo dự án STS"],
              ],
              widths=[1.2, 1.4, 1.6, 2.2])

    add_page_number_footer(doc)

    out = os.path.join(os.path.dirname(__file__), "bao-cao-du-an-sts.docx")
    doc.save(out)
    print("Đã tạo:", out)


if __name__ == "__main__":
    build()
